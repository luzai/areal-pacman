from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from areal_pacman.synthetic.env import PacmanEnv
from areal_pacman.synthetic.maze_suite import maze_records, suite_summary
from areal_pacman.synthetic.vision import render_env_image


RUN_ROOT = ROOT / "run_artifacts" / "multi_maze_v1_20260715"
SAFE_RUN_ROOT = ROOT / "run_artifacts" / "multi_maze_v1_safeprogress_alpha4_20260715"
ASSET_ROOT = RUN_ROOT / "report_assets"
REPORT_ROOT = ROOT.parent / "reports"
OUT = REPORT_ROOT / "multi_maze_vlm_rl_report_2026-07-15.pptx"

BG = RGBColor(247, 248, 250)
INK = RGBColor(25, 31, 42)
MUTED = RGBColor(92, 102, 119)
BLUE = RGBColor(42, 88, 184)
GREEN = RGBColor(31, 130, 102)
RED = RGBColor(202, 72, 72)
AMBER = RGBColor(198, 132, 36)
WHITE = RGBColor(255, 255, 255)


def add_text(slide, text: str, x: float, y: float, w: float, h: float, size: int = 18,
             bold: bool = False, color: RGBColor = INK, align=PP_ALIGN.LEFT) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Aptos"


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    add_text(slide, title, 0.55, 0.28, 12.2, 0.52, 25, True)
    if subtitle:
        add_text(slide, subtitle, 0.58, 0.84, 12.0, 0.35, 11, False, MUTED)


def add_footer(slide) -> None:
    add_text(
        slide,
        "vla-research/areal-pacman | artifact-backed | 2026-07-15",
        0.55,
        7.08,
        12.2,
        0.2,
        8,
        False,
        MUTED,
        PP_ALIGN.RIGHT,
    )


def add_bullets(slide, items: list[str], x: float, y: float, w: float, h: float,
                size: int = 16, color: RGBColor = INK) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.name = "Aptos"
        p.space_after = Pt(8)
        p.level = 0


def add_panel(slide, x: float, y: float, w: float, h: float, fill: RGBColor = WHITE,
              line: RGBColor = RGBColor(205, 210, 220)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def add_picture(slide, path: Path, x: float, y: float, w: float, h: float | None = None) -> None:
    kwargs = {"width": Inches(w)}
    if h is not None:
        kwargs["height"] = Inches(h)
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kwargs)


def load_comparison(root: Path) -> dict[str, object]:
    return json.loads((root / "comparison.json").read_text(encoding="utf-8"))


def make_maze_grid() -> Path:
    ASSET_ROOT.mkdir(parents=True, exist_ok=True)
    samples = []
    for split in ("train", "validation", "test"):
        records = maze_records(split)
        for record in (records[0], records[len(records) // 2], records[-1]):
            env = PacmanEnv(layout_name=str(record["name"]), max_steps=30)
            samples.append((split, str(record["name"]), render_env_image(env, tile_size=32)))

    cell_w, cell_h = 340, 260
    canvas = Image.new("RGB", (cell_w * 3, cell_h * 3), (247, 248, 250))
    draw = ImageDraw.Draw(canvas)
    label_font = ImageFont.load_default()
    colors = {"train": (42, 88, 184), "validation": (198, 132, 36), "test": (31, 130, 102)}
    for index, (split, name, image) in enumerate(samples):
        col, row = index % 3, index // 3
        x, y = col * cell_w, row * cell_h
        draw.rectangle((x + 8, y + 8, x + cell_w - 8, y + cell_h - 8), fill=(255, 255, 255), outline=(205, 210, 220), width=2)
        rendered = image.resize((288, 192), Image.Resampling.NEAREST)
        canvas.paste(rendered, (x + 26, y + 38))
        draw.rectangle((x + 8, y + 8, x + cell_w - 8, y + 30), fill=colors[split])
        draw.text((x + 18, y + 13), f"{split.upper()} | {name}", fill=(255, 255, 255), font=label_font)
    output = ASSET_ROOT / "maze_split_grid.png"
    canvas.save(output)
    return output


def make_result_chart(comparison: dict[str, object], name: str) -> Path:
    width, height = 1400, 680
    image = Image.new("RGB", (width, height), (247, 248, 250))
    draw = ImageDraw.Draw(image)
    regular = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 26)
    bold = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 30)
    small = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 22)
    splits = ("train", "validation", "test")
    max_rate = 0.15
    left, top, chart_h = 120, 95, 430
    draw.text((left, 25), "Macro pass rate: each maze has equal weight", fill=(25, 31, 42), font=bold)
    for tick in range(4):
        rate = tick * 0.05
        y = top + chart_h - rate / max_rate * chart_h
        draw.line((left, y, width - 70, y), fill=(215, 219, 227), width=2)
        draw.text((35, y - 14), f"{rate:.0%}", fill=(92, 102, 119), font=small)
    group_w = 350
    bar_w = 96
    for index, split in enumerate(splits):
        values = comparison["splits"][split]
        base = float(values["baseline"]["macro_layout_pass_rate"])
        post = float(values["post_rl"]["macro_layout_pass_rate"])
        center = left + 220 + index * group_w
        for offset, value, color, label in ((-58, base, (42, 88, 184), "Base"), (58, post, (31, 130, 102), "Post")):
            x0 = center + offset - bar_w // 2
            y0 = top + chart_h - value / max_rate * chart_h
            draw.rectangle((x0, y0, x0 + bar_w, top + chart_h), fill=color)
            draw.text((x0 + 10, y0 - 34), f"{value:.1%}", fill=(25, 31, 42), font=small)
            draw.text((x0 + 16, top + chart_h + 12), label, fill=(92, 102, 119), font=small)
        draw.text((center - 70, top + chart_h + 60), split.upper(), fill=(25, 31, 42), font=regular)
        delta = post - base
        draw.text((center - 68, top + chart_h + 102), f"delta {delta:+.1%}", fill=(31, 130, 102) if delta >= 0 else (202, 72, 72), font=small)
    output = ASSET_ROOT / name
    image.save(output)
    return output


def add_metric(slide, value: str, label: str, x: float, y: float, color: RGBColor) -> None:
    add_panel(slide, x, y, 2.15, 1.05)
    box = slide.shapes.add_textbox(Inches(x + 0.08), Inches(y + 0.16), Inches(1.99), Inches(0.34))
    tf = box.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = value
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(17 if len(value) > 8 else 24)
    p.font.bold = True
    p.font.color.rgb = color
    p.font.name = "Aptos"
    add_text(slide, label, x + 0.12, y + 0.64, 1.91, 0.24, 10, False, MUTED, PP_ALIGN.CENTER)


def build(output: Path = OUT, static_demos: bool = False) -> None:
    sparse_comparison = load_comparison(RUN_ROOT)
    safe_comparison = load_comparison(SAFE_RUN_ROOT)
    summary = suite_summary()
    maze_grid = make_maze_grid()
    sparse_chart = make_result_chart(sparse_comparison, "macro_pass_rate_sparse_chart.png")
    safe_chart = make_result_chart(safe_comparison, "macro_pass_rate_safe_progress_chart.png")
    output.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Multi-Maze Image-Only VLM RL", "Qwen3.5-9B + AReaL | 64 train / 16 validation / 32 held-out test mazes")
    add_text(slide, "Safe progress passed. Generalization did not.", 0.7, 1.35, 5.8, 0.7, 25, True, RED)
    add_bullets(slide, [
        "Alpha-4 safe-progress GRPO completed 24/24 optimizer steps.",
        "Validation selected epoch 0 from 12.9%, 10.1%, and 10.6%.",
        "Train: 3.9% -> 5.1%; validation: 12.5% -> 12.9%.",
        "Held-out test: 5.1% -> 0.0%. No evidence of transfer.",
    ], 0.75, 2.2, 5.65, 3.5, 18)
    add_picture(slide, safe_chart, 6.65, 1.45, 6.0)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Generalization Protocol", "Different wall topology, not merely different random seeds on one maze")
    add_picture(slide, maze_grid, 0.55, 1.25, 7.3)
    add_bullets(slide, [
        "112 deterministic layouts: 64 train, 16 validation, 32 test.",
        "112 unique full-layout hashes and 112 unique wall-topology hashes.",
        "Zero topology overlap across splits; five pellets per maze.",
        f"Oracle replay passes every maze in {summary['oracle_steps_min']}..{summary['oracle_steps_max']} steps, under max_steps=30.",
        "Primary metric is macro per-layout pass rate, so asynchronous oversampling cannot dominate the result.",
    ], 8.1, 1.45, 4.65, 4.9, 16)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Exact Image-Only Model Contract", "Only the changing state is pixels; fixed game rules remain in the system prompt")
    nodes = [
        ("ENV", "PacMan state", 0.55, BLUE),
        ("RENDER", "one 288 x 192 PNG", 2.55, BLUE),
        ("VLM CALL", "system rules + user text + PNG", 4.75, GREEN),
        ("CHOICE", "legal non-backtracking action", 7.15, AMBER),
        ("ENV STEP", "reward + terminal", 9.55, BLUE),
        ("TRACE", "prompt + response + result", 11.45, RED),
    ]
    for title, body, x, color in nodes:
        add_panel(slide, x, 1.45, 1.55 if x != 4.75 else 1.9, 1.15)
        add_text(slide, title, x + 0.05, 1.62, 1.45 if x != 4.75 else 1.8, 0.28, 12, True, color, PP_ALIGN.CENTER)
        add_text(slide, body, x + 0.08, 2.0, 1.39 if x != 4.75 else 1.74, 0.42, 10, False, MUTED, PP_ALIGN.CENTER)
    for x in (2.15, 4.35, 6.82, 9.18, 11.08):
        arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(1.82), Inches(0.3), Inches(0.28))
        arrow.fill.solid(); arrow.fill.fore_color.rgb = MUTED; arrow.line.fill.background()
    add_panel(slide, 0.7, 3.05, 5.95, 2.9, RGBColor(255, 255, 255))
    add_text(slide, "Fixed system prompt", 0.95, 3.28, 5.4, 0.35, 17, True, BLUE)
    add_bullets(slide, [
        "Visual legend: blue walls, dark floor, yellow PacMan, red ghost, pale-yellow pellets.",
        "Ghost alternates stay and move turns; movement activates only when Manhattan distance <= 3.",
        "Scoring, collision, pellet-clearing win condition, max-step stop, and one-token output rule.",
    ], 0.95, 3.78, 5.4, 1.85, 14)
    add_panel(slide, 6.95, 3.05, 5.65, 2.9, RGBColor(255, 255, 255))
    add_text(slide, "Per-turn user message", 7.2, 3.28, 5.1, 0.35, 17, True, GREEN)
    add_text(slide, '"Choose exactly one action token: up, down, left, right, or stay."', 7.2, 3.85, 5.0, 0.75, 16, True)
    add_bullets(slide, [
        "+ exactly one current PNG frame.",
        "No ASCII maze, coordinates, score, pellet count, route hint, or frame history.",
    ], 7.2, 4.72, 5.0, 1.0, 14)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Sparse GRPO Control", "Final epoch checkpoint; macro pass rate gives every maze equal weight")
    add_picture(slide, sparse_chart, 0.55, 1.22, 7.0)
    add_metric(slide, "3.9% -> 5.5%", "TRAIN", 8.0, 1.5, GREEN)
    add_metric(slide, "12.5% -> 4.7%", "VALIDATION", 10.35, 1.5, RED)
    add_metric(slide, "5.1% -> 3.1%", "HELD-OUT TEST", 8.0, 2.8, RED)
    add_metric(slide, "0 / 0", "PARSE / ILLEGAL STEPS", 10.35, 2.8, GREEN)
    add_bullets(slide, [
        "Raw test: 7/132 before RL, 5/137 after RL.",
        "Layouts with any test win: 7/32 before, 3/32 after.",
        "Validation-selected checkpointing was not used; this run evaluates the final epoch-3 checkpoint.",
        "Result is constrained VLM control, not unconstrained raw text generation.",
    ], 8.0, 4.2, 4.65, 2.0, 14)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Safe-Progress Reward Ablation", "Hidden env state shapes reward; model input remains one current PNG")
    add_panel(slide, 0.65, 1.28, 5.4, 1.52)
    add_text(slide, "reward = original + alpha *\n(d_before - d_after)", 0.9, 1.5, 4.9, 0.7, 17, True, BLUE, PP_ALIGN.CENTER)
    add_text(slide, "alpha = 4; collision-free BFS under the real ghost phase", 0.9, 2.35, 4.9, 0.25, 10, False, MUTED, PP_ALIGN.CENTER)
    add_picture(slide, safe_chart, 6.25, 1.2, 6.45)
    add_bullets(slide, [
        "Epoch validation macro: 12.91%, 10.07%, 10.63%; epoch 0 selected before test.",
        "Selected train: 3.91% -> 5.08%; validation: 12.50% -> 12.91%.",
        "Held-out test: 5.08% -> 0.00% (0/130 trajectories, 0/32 layouts with any win).",
        "Reward audit: 18,032 training steps, zero alpha or formula mismatches.",
        "Conclusion: denser progress signal alone did not improve unseen-topology control.",
    ], 0.85, 3.05, 5.3, 3.2, 15)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Training And Checkpoint Evidence", "The optimization system ran correctly; the learned policy did not generalize")
    add_metric(slide, "24 / 24", "GRPO TRAIN STEPS", 0.7, 1.45, GREEN)
    add_metric(slide, "3", "TRAIN EPOCHS", 3.1, 1.45, BLUE)
    add_metric(slide, "~46 min", "8xH100 TRAIN TIME", 5.5, 1.45, BLUE)
    add_metric(slide, "1e-6", "LEARNING RATE", 7.9, 1.45, AMBER)
    add_metric(slide, "768", "TARGET TRAIN ROLLOUTS", 10.3, 1.45, BLUE)
    add_bullets(slide, [
        "Sparse and alpha-4 safe-progress runs both use four samples, legal non-backtracking decode, and 24 updates.",
        "Each alpha-4 checkpoint has 427 trained keys plus 333 restored frozen visual keys in a 912 MB shard.",
        "Across six safe-progress summaries: zero parse failures and zero illegal actions.",
        "Prompt audit: 36,043 model calls, each with exact fixed text plus one PNG and no dynamic maze text.",
        "Evidence rejects both short GRPO recipes as unseen-maze generalization improvements.",
    ], 0.85, 3.1, 11.75, 3.25, 17)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Sparse-Control Test Exemplars", "Alpha-4 safe progress had no test success; these preserve the earlier sparse examples")
    success = ASSET_ROOT / ("post_rl_test_success_contact_sheet.png" if static_demos else "post_rl_test_success.gif")
    failure = ASSET_ROOT / ("post_rl_test_failure_contact_sheet.png" if static_demos else "post_rl_test_failure.gif")
    add_picture(slide, success, 0.55, 1.35, 6.0, 4.65)
    add_picture(slide, failure, 6.78, 1.35, 6.0, 4.65)
    add_text(slide, "SUCCESS | mmv1_test_005 | 19 steps | all pellets", 0.7, 6.15, 5.7, 0.32, 13, True, GREEN, PP_ALIGN.CENTER)
    add_text(slide, "FAILURE | mmv1_test_007 | 5 steps | caught", 6.95, 6.15, 5.65, 0.32, 13, True, RED, PP_ALIGN.CENTER)
    add_text(slide, "GIF terminal frames are retained for 2600 ms. PDF uses full contact sheets.", 2.4, 6.62, 8.55, 0.3, 11, False, MUTED, PP_ALIGN.CENTER)
    add_footer(slide)

    prs.save(output)
    print(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=OUT)
    parser.add_argument("--static-demos", action="store_true")
    args = parser.parse_args()
    build(args.output, static_demos=args.static_demos)


if __name__ == "__main__":
    main()
