from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT.parent / "reports" / "vision_side_skeleton_report_2026-07-14.pptx"
IMAGEONLY_NOTRAIN_DIR = (
    ROOT / "run_artifacts" / "areal_vlm_imageonly_rules_activationgate_notrain16_20260713"
)
CURRENT_MAZE_IMAGE = IMAGEONLY_NOTRAIN_DIR / "current_maze_initial_frame.png"


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12.4), Inches(0.55))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(25)
    p.font.bold = True
    p.font.name = "Aptos Display"
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.48), Inches(0.82), Inches(12.0), Inches(0.35))
        sp = sub_box.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(11)
        sp.font.color.rgb = RGBColor(80, 88, 104)


def add_bullets(slide, items: list[str], x: float, y: float, w: float, h: float, size: int = 15) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.space_after = Pt(6)


def add_footer(slide) -> None:
    box = slide.shapes.add_textbox(Inches(0.45), Inches(7.05), Inches(12.4), Inches(0.25))
    p = box.text_frame.paragraphs[0]
    p.text = "vla-research/areal-pacman | generated from local artifacts | 2026-07-14"
    p.font.size = Pt(8)
    p.alignment = PP_ALIGN.RIGHT


def add_image(slide, path: Path, x: float, y: float, w: float | None = None, h: float | None = None) -> None:
    if not path.exists():
        return
    kwargs = {}
    if w is not None:
        kwargs["width"] = Inches(w)
    if h is not None:
        kwargs["height"] = Inches(h)
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kwargs)


def add_flow_node(
    slide,
    title: str,
    body: str,
    x: float,
    y: float,
    w: float,
    h: float,
    accent: RGBColor,
) -> None:
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(248, 249, 251)
    panel.line.color.rgb = RGBColor(196, 202, 212)

    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.42))
    header.fill.solid()
    header.fill.fore_color.rgb = accent
    header.line.color.rgb = accent
    hp = header.text_frame.paragraphs[0]
    hp.text = title
    hp.font.size = Pt(12)
    hp.font.bold = True
    hp.font.color.rgb = RGBColor(255, 255, 255)
    hp.alignment = PP_ALIGN.CENTER

    box = slide.shapes.add_textbox(Inches(x + 0.12), Inches(y + 0.52), Inches(w - 0.24), Inches(h - 0.62))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = body
    p.font.size = Pt(11)
    p.alignment = PP_ALIGN.CENTER


def add_right_arrow(slide, x: float, y: float, w: float = 0.38, h: float = 0.28) -> None:
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RGBColor(118, 126, 140)
    arrow.line.fill.background()


def add_merge_line(slide, x1: float, y1: float, x2: float, y2: float) -> None:
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = RGBColor(118, 126, 140)
    line.line.width = Pt(1.5)


def build(output: Path = OUT, static_demos: bool = False) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Image-Only VLM PacMan RL", "AReaL same-maze overfit from a corrected no-train baseline")
    add_bullets(
        slide,
        [
            "Qwen3.5-9B receives one rendered PNG per turn and returns one action token.",
            "Corrected no-train baseline: 8 wins / 32 rollouts = 25.0%.",
            "Six-epoch GRPO run completed all 24 train steps from Qwen3.5-9B.",
            "Fresh final-checkpoint evaluation: 24 wins / 36 rollouts = 66.7%.",
            "Absolute gain: +41.7 points; pass-rate ratio: 2.67x.",
            "This is constrained inference with legal non-backtracking choices, not raw generation.",
        ],
        0.65,
        1.35,
        6.1,
        4.8,
        size=19,
    )
    add_image(slide, CURRENT_MAZE_IMAGE, 7.3, 1.45, w=4.8)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "What The VLM Receives", "Frame-by-frame control, not a video-file input")
    add_bullets(
        slide,
        [
            "System prompt: visual legend, simplified rules, scoring, win/stop conditions, and action-only output format.",
            "Visual legend: blue walls, dark floor, yellow PacMan, red ghost, pale yellow pellets.",
            "Ghost rule: alternate inactive/movement turns; on a movement turn, activate only at Manhattan distance <= 3.",
            "When active, the ghost takes one legal tile minimizing Manhattan distance; at distance > 3 it stays still.",
            "User each turn: 'Choose exactly one action token: up, down, left, right, or stay.' plus the current PNG.",
            "No dynamic text state: no ASCII maze, coordinates, positions, score, pellet count, or route hint.",
            "Image-only means the changing maze state is pixels only; fixed rules remain text.",
            "Decode control: structured choice exposes only legal, non-backtracking action tokens.",
            "Loop: PNG frame -> VLM -> action -> env step -> next PNG -> trajectory.",
        ],
        0.55,
        1.05,
        7.1,
        5.8,
        size=15,
    )
    add_image(slide, CURRENT_MAZE_IMAGE, 8.0, 1.35, w=4.7)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Pure Text To Image-Only VLM", "One shared closed loop, with two observation and model-call adapters")
    add_flow_node(slide, "ENV STATE", "PacMan, ghost, pellets", 0.35, 2.62, 1.35, 1.35, RGBColor(70, 77, 90))

    add_merge_line(slide, 1.70, 3.05, 2.15, 2.28)
    add_merge_line(slide, 1.70, 3.55, 2.15, 4.48)
    add_flow_node(
        slide,
        "PURE-TEXT OBS",
        "ASCII grid + textual state metadata",
        2.15,
        1.58,
        2.05,
        1.42,
        RGBColor(41, 98, 146),
    )
    add_flow_node(
        slide,
        "IMAGE-ONLY OBS",
        "Current PNG + fixed action sentence",
        2.15,
        3.78,
        2.05,
        1.42,
        RGBColor(31, 122, 92),
    )

    add_right_arrow(slide, 4.30, 2.14)
    add_right_arrow(slide, 4.30, 4.34)
    add_flow_node(
        slide,
        "STRING MODEL API",
        "completion, chat_user, or chat",
        4.80,
        1.58,
        1.85,
        1.42,
        RGBColor(41, 98, 146),
    )
    add_flow_node(
        slide,
        "MULTIMODAL CHAT",
        "System rules + user text + image_url",
        4.80,
        3.78,
        1.85,
        1.42,
        RGBColor(31, 122, 92),
    )

    add_merge_line(slide, 6.65, 2.28, 7.15, 3.02)
    add_merge_line(slide, 6.65, 4.48, 7.15, 3.62)
    add_flow_node(
        slide,
        "CONSTRAINED ACTION",
        "Legal choice + no immediate backtracking",
        7.15,
        2.48,
        1.75,
        1.68,
        RGBColor(102, 75, 140),
    )
    add_right_arrow(slide, 9.00, 3.18)
    add_flow_node(
        slide,
        "ENV STEP",
        "Transition + reward + terminal check",
        9.48,
        2.48,
        1.62,
        1.68,
        RGBColor(70, 77, 90),
    )
    add_right_arrow(slide, 11.20, 3.18)
    add_flow_node(
        slide,
        "TRACE JSON",
        "Prompt + response + action + result",
        11.68,
        2.48,
        1.30,
        1.68,
        RGBColor(147, 78, 48),
    )

    add_flow_node(
        slide,
        "POST-HOC REPLAY",
        "GIF/PDF for humans only. Never model input.",
        10.55,
        4.72,
        2.43,
        1.15,
        RGBColor(147, 78, 48),
    )
    down_arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(12.05), Inches(4.20), Inches(0.28), Inches(0.38))
    down_arrow.fill.solid()
    down_arrow.fill.fore_color.rgb = RGBColor(118, 126, 140)
    down_arrow.line.fill.background()

    loop_arrow = slide.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, Inches(0.52), Inches(6.12), Inches(11.82), Inches(0.38))
    loop_arrow.fill.solid()
    loop_arrow.fill.fore_color.rgb = RGBColor(224, 228, 235)
    loop_arrow.line.color.rgb = RGBColor(174, 181, 193)
    lp = loop_arrow.text_frame.paragraphs[0]
    lp.text = "NEXT TURN: render the new current observation"
    lp.font.size = Pt(12)
    lp.font.bold = True
    lp.font.color.rgb = RGBColor(65, 72, 84)
    lp.alignment = PP_ALIGN.CENTER
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "No-Train vs Post-RL", "Fresh no-update evaluation on the same maze and prompt contract")
    add_bullets(
        slide,
        [
            "Pre-RL: 8/32 wins = 25.0%; reasons 8 all_pellets, 4 caught, 20 max_steps.",
            "Post-RL: 24/36 wins = 66.7%; reasons 24 all_pellets, 12 caught, 0 max_steps.",
            "Post-RL reward min / max / mean: -39.0 / 140.0 / 80.56.",
            "Post-RL evaluator used lr=0 for all 4 steps: no additional model update.",
            "420/420 post-RL calls passed the corrected activation-gate prompt audit.",
            "Post-RL parse failures: 0. Illegal actions after structured choice: 0.",
            "Both runs repeat one medium_default maze; episode ids are sampling labels.",
            "Conclusion: clear same-maze partial overfit, not saturated 100% performance.",
        ],
        0.8,
        1.25,
        11.7,
        5.4,
        size=19,
    )
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    success_subtitle = (
        "Static PDF contact sheet | final terminal frame included"
        if static_demos
        else "Animated GIF | final post-action terminal state included"
    )
    add_title(slide, "No-Train Success Exemplar", success_subtitle)
    add_bullets(
        slide,
        [
            "Episode pacman-episode-3: WIN, all pellets, 10 steps, reward 140.0.",
            "Actions: right, right, down, down, right, right, up, up, right, right.",
            "Frames 1-10 are exact model observations before each action.",
            "Final frame is the env state after action 10: all pellets cleared.",
            "Full JSON records system prompt, fixed user text, image hash, response, and env result.",
        ],
        0.45,
        1.0,
        4.0,
        5.8,
        size=15,
    )
    success_media = "no_train_success_demo_contact_sheet.png" if static_demos else "no_train_success_demo.gif"
    add_image(slide, IMAGEONLY_NOTRAIN_DIR / success_media, 5.25, 1.15, h=5.65)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    failure_subtitle = (
        "Static PDF contact sheet | final terminal frame included"
        if static_demos
        else "Animated GIF | final post-action terminal state included"
    )
    add_title(slide, "No-Train Failure Exemplar", failure_subtitle)
    add_bullets(
        slide,
        [
            "Episode pacman-episode-11: FAIL, caught, 13 steps, reward -33.0.",
            "Actions: down, down, down, right, right, up, right, right, up, up, right, right, down.",
            "Frames 1-13 are exact model observations before each action.",
            "Final frame is the env state after action 13: PacMan and ghost collide.",
            "Different outcome comes from stochastic sampling under legal non-backtracking choices.",
        ],
        0.45,
        1.0,
        4.0,
        5.8,
        size=15,
    )
    failure_media = "no_train_failure_demo_contact_sheet.png" if static_demos else "no_train_failure_demo.gif"
    add_image(slide, IMAGEONLY_NOTRAIN_DIR / failure_media, 5.25, 1.15, h=5.65)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Conclusion", "What the completed RL gate does and does not establish")
    add_bullets(
        slide,
        [
            "Established: AReaL can execute the complete image -> VLM -> action -> env -> trajectory loop.",
            "Established: six-epoch image-only VLM RL completed from the corrected Qwen3.5-9B baseline.",
            "Established: fresh no-update pass rate improved from 25.0% to 66.7% on this one maze.",
            "Interpretation: clear partial overfit, but not near-100% saturated memorization.",
            "Not established: raw unconstrained VLM pass rate.",
            "Not established: generalization to different maze layouts or initial states.",
            "Checkpoint fix: restored 333 frozen visual tensors without replacing 427 trained weights.",
        ],
        0.8,
        1.35,
        11.7,
        4.8,
        size=21,
    )
    add_footer(slide)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build the compact vision-side PacMan report deck.")
    parser.add_argument("--output", type=Path, default=OUT)
    parser.add_argument("--static-demos", action="store_true", help="Use terminal PNGs for PDF conversion.")
    args = parser.parse_args()
    build(output=args.output, static_demos=args.static_demos)
